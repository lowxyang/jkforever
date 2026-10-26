#!/usr/bin/env python3
"""Serve the invitation and append each RSVP to RSVPs.xlsx."""

from datetime import datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import json

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "RSVPs.xlsx"
HEADERS = ["Timestamp", "Attending", "Attendee(s)", "Dietary requirements", "Email / WhatsApp"]
WIDTHS = [22, 22, 36, 32, 36]


class InvitationHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def end_headers(self):
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def do_POST(self):
        if self.path.split("?", 1)[0] != "/api/rsvp":
            self.send_error(404)
            return
        length = int(self.headers.get("Content-Length", "0") or 0)
        raw = self.rfile.read(length)
        try:
            data = json.loads(raw.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            self.respond(400, {"ok": False, "error": "Invalid JSON"})
            return
        if not isinstance(data, dict):
            self.respond(400, {"ok": False, "error": "Invalid RSVP"})
            return

        attendees = str(data.get("attendees") or data.get("name") or "").strip()
        attending = str(data.get("attending", "")).strip().lower()
        if not attendees or attending not in {"yes", "no"}:
            self.respond(400, {"ok": False, "error": "Attendee name(s) and attendance are required."})
            return

        row = [
            str(data.get("submittedAt") or datetime.now(timezone.utc).isoformat()),
            "Joyfully accept" if attending == "yes" else "Regretfully decline",
            attendees,
            str(data.get("diet", "")).strip(),
            str(data.get("contact") or data.get("email") or "").strip(),
        ]
        try:
            append_rsvp(row)
        except OSError:
            self.respond(500, {"ok": False, "error": "Could not write RSVPs.xlsx. Close the file if it is open in Excel."})
            return
        self.respond(200, {"ok": True})

    def respond(self, status, payload):
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def has_current_layout(path):
    sheet = load_workbook(path).active
    return [sheet.cell(1, index).value for index in range(1, len(HEADERS) + 1)] == HEADERS


def append_rsvp(row):
    if WORKBOOK.exists() and not has_current_layout(WORKBOOK):
        # The questions changed: keep the old file and start a fresh one.
        stamp = datetime.now().strftime("%Y-%m-%d %H-%M")
        WORKBOOK.rename(WORKBOOK.with_name(f"RSVPs (old layout {stamp}).xlsx"))
    if WORKBOOK.exists():
        workbook = load_workbook(WORKBOOK)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "RSVPs"
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        for index, width in enumerate(WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.append(row)
    workbook.save(WORKBOOK)


def main():
    server = ThreadingHTTPServer(("127.0.0.1", 8080), InvitationHandler)
    print("Invitation at http://127.0.0.1:8080")
    print(f"RSVPs will be added to {WORKBOOK}")
    server.serve_forever()


if __name__ == "__main__":
    main()
